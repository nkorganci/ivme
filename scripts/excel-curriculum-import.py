"""Import ONLY worksheet zero into the existing shared category registry."""
import argparse
import hashlib
import json
import subprocess
from pathlib import Path

import openpyxl

APP = Path(__file__).resolve().parents[1]
SOURCE = APP.parent / 'Kaynaklar/mufredat/YKS_TYT_AYT_Konu_Soru_Dagilimi_2018-2026.xlsx'


def first_sheet(path=SOURCE):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        header = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
        if tuple(header[:6]) != ('Ders', 'TYT', 'AYT', 'Sınıf', 'Ünite', 'Konu'):
            raise ValueError('İlk sayfanın beklenen başlıkları bulunamadı.')
        rows = []
        for number, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
            if not any(row):
                continue
            if not all(row[i] for i in (0, 3, 4, 5)) or not (row[1] or row[2]):
                raise ValueError(f'Geçersiz konu satırı: {number}')
            counts = list(row[6:10]) + list(row[12:30])
            if any(v is not None and (not isinstance(v, int) or isinstance(v, bool) or v < 0) for v in counts):
                raise ValueError(f'Geçersiz soru sayısı: {number}')
            rows.append(dict(row=number, subject=row[0], tyt=row[1], ayt=row[2], gradeLabel=str(row[3]),
                unit=row[4], topic=row[5], tytTotal=row[6], aytTotal=row[7], total=row[8], recent=row[9],
                verification=row[10], note=row[11], tytYears=list(row[12:21]), aytYears=list(row[21:30]),
                tytKey=row[30], aytKey=row[31]))
        encoded = json.dumps(rows, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
        return dict(file=path.relative_to(APP.parent).as_posix(), sheet=sheet.title, sheetIndex=0,
            sha256=hashlib.sha256(encoded).hexdigest(), years=list(range(2018, 2027)), rows=rows)
    finally:
        book.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    subprocess.run(['node', 'scripts/excel-curriculum-import.mjs', *(['--check'] if args.check else [])],
        input=json.dumps(first_sheet(), ensure_ascii=False).encode('utf-8'), cwd=APP, check=True)
