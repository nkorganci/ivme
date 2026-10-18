"""Export an editable physics taxonomy proposal, or validate a teacher-edited workbook.

Only reads PostgreSQL. Never publishes labels or writes teacher/student records.
The workbook is a review artifact of the existing registry, not another dictionary.
"""
import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from kaynakli_arastirma_db import db_connection, query

APP = Path(__file__).resolve().parents[1]
TAXONOMY = APP / 'backend/src/main/resources/content-taxonomy.json'
STATE = APP / 'veri/content-labels/taxonomy-review'
TOPIC_HEADERS = ['Kod', 'Sınav', 'Temel konu adı', 'Kapsam / öğretmen notu', 'Öğretmen kararı']
MAP_HEADERS = ['Eski kategori kimliği', 'Sınav', 'Eski konu', 'Eski alt başlık',
               'Önerilen temel konu kodu', 'Önerinin dayanağı', 'Öğretmen kararı',
               'Öğretmen notu', 'Soru', 'PDF', 'Video', 'Diğer', 'Örnek kaynak kimlikleri']


def digest(data):
    return hashlib.sha256(data).hexdigest()


def normal(value):
    text = str(value or '').lower().replace('ı', 'i')
    text = ''.join(c for c in unicodedata.normalize('NFD', text) if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', ' ', text).strip()


def read_snapshot():
    # One read-only transaction prevents mixing two published catalog runs.
    sql = """BEGIN TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY;
    SELECT json_build_object(
      'runKey',(SELECT run_key FROM learning_catalog.current_run),
      'categories',(SELECT json_agg(definition ORDER BY category_id)
         FROM learning_catalog.current_categories WHERE definition->>'subjectId'='fizik'),
      'impact',(SELECT json_agg(row_to_json(x)) FROM (
         SELECT a.category_id,i.kind,count(DISTINCT i.item_key) AS count,
           (array_agg(DISTINCT i.item_key ORDER BY i.item_key))[1:3] AS examples
         FROM learning_catalog.current_assignments a
         JOIN learning_catalog.current_categories c ON c.category_id=a.category_id
         JOIN learning_catalog.current_items i ON i.item_key=a.item_key
         WHERE c.definition->>'subjectId'='fizik' GROUP BY a.category_id,i.kind) x),
      'uniqueResources',(SELECT json_object_agg(kind,n) FROM (
         SELECT i.kind,count(DISTINCT i.item_key) n
         FROM learning_catalog.current_items i
         JOIN learning_catalog.current_assignments a ON a.item_key=i.item_key
         JOIN learning_catalog.current_categories c ON c.category_id=a.category_id
         WHERE c.definition->>'subjectId'='fizik' GROUP BY i.kind) x));
    COMMIT;"""
    result = json.loads(query(db_connection(), sql))
    if not result['runKey'] or not result['categories']:
        raise ValueError('Yayımlanmış fizik kataloğu bulunamadı.')
    registry = json.loads(TAXONOMY.read_text(encoding='utf-8'))
    definitions = {c['id']: c for c in registry['categories']}
    if any(definitions.get(c['id']) != c for c in result['categories']):
        raise ValueError('DB ve bakım sözlüğü farklı; inceleme öncesi sürümleri uzlaştırın.')
    result['taxonomySha256'] = digest(TAXONOMY.read_bytes())
    return result


def proposal(snapshot):
    categories = snapshot['categories']
    by_id = {c['id']: c for c in categories}
    roots = sorted((c for c in categories if c.get('origin') == 'CURRICULUM'
                    and not c.get('subtopic')), key=lambda c: (c['exam'], normal(c['topic'])))
    codes = {c['id']: f'F{i:03}' for i, c in enumerate(roots, 1)}
    terms = defaultdict(set)
    for c in roots:
        for title in [c['topic'], *c.get('aliases', [])]:
            terms[(c['exam'], normal(title))].add(c['id'])
    impacts = defaultdict(Counter)
    examples = defaultdict(list)
    for row in snapshot['impact'] or []:
        impacts[row['category_id']][row['kind']] = row['count']
        examples[row['category_id']].extend(row['examples'])
    mappings = []
    for c in categories:
        target, reason = '', 'ESLESME_BEKLIYOR'
        if c['id'] in codes:
            target, reason = c['id'], 'MEVCUT_TEMEL_KONU'
        elif c.get('parentId'):
            parent, visited = c, set()
            while parent.get('parentId') and parent['id'] not in visited:
                visited.add(parent['id'])
                parent = by_id.get(parent['parentId'], {})
            if parent.get('id') in codes and parent.get('exam') == c['exam']:
                target, reason = parent['id'], 'MEVCUT_UST_KONU_ONERISI'
        if not target:
            matches = terms[(c['exam'], normal(c['topic']))]
            if len(matches) == 1:
                target, reason = next(iter(matches)), 'BIREBIR_AD_VEYA_KAYITLI_ALIAS'
        counts = impacts[c['id']]
        mappings.append([c['id'], c['exam'], c['topic'], c.get('subtopic', ''),
                         codes.get(target, ''), reason, 'BEKLIYOR', '',
                         counts['QUESTION'], counts['PDF'], counts['VIDEO'],
                         sum(n for k, n in counts.items() if k not in {'QUESTION', 'PDF', 'VIDEO'}),
                         ' | '.join(examples[c['id']][:4])])
    topics = [[codes[c['id']], c['exam'], c['topic'], '', 'BEKLIYOR'] for c in roots]
    return topics, mappings


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def export():
    snapshot = read_snapshot()
    topics, mappings = proposal(snapshot)
    directory = STATE / datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    directory.mkdir(parents=True, exist_ok=False)
    snapshot.update({'version': 1, 'createdAt': datetime.now(timezone.utc).isoformat(),
                     'topicRows': topics, 'mappingRows': mappings})
    write_json(directory / 'baseline.json', snapshot)
    workbook = Workbook()
    info = workbook.active
    info.title = 'Başlangıç'
    lines = [
        ['Fizik temel konu incelemesi', 'Öğretmen taslağı; canlı veriye uygulanmadı.'],
        ['Kaynak sürümü', snapshot['runKey']],
        ['Mevcut temel konu', len(topics)], ['İncelenecek eski kategori', len(mappings)],
        ['1', 'Temel konular sayfasında adları/kapsamı düzenleyin. Yeni konu için F038 gibi benzersiz kod ekleyin.'],
        ['2', '46 sayısı sizin hedefinizdir; mevcut 37 başlık onaylı liste değildir. TYT/AYT ayrı kapsamdır.'],
        ['3', 'Konuları birleştirirken kaldırdığınız kodu eşleştirmelerde de kalan koda çevirin.'],
        ['4', 'Eşleştirmeler sayfasında önerilen kodu düzeltin. Boş önerilerde hedefi siz seçin.'],
        ['5', 'Kesin eşleştirmeye ONAY yazın. Geniş/karma/yanlış konudaki kayda ICERIK_INCELE yazın ve not ekleyin.'],
        ['6', 'ICERIK_INCELE için hedef kodu boşaltın; her kaynak ayrıca incelenir. Otomatik çoklu konu verilmez.'],
        ['7', 'Üst konu önerisi anlam doğruluğu veya içerik onayı değildir. Gerekirse bu öneriyi değiştirin.'],
        ['8', 'Temel konular için de ONAY seçin. Eski kimlik/ad, sayım ve dayanak sütunlarını değiştirmeyin.'],
        ['9', 'Kaydedilmiş XLSX dosyasını baseline.json ile aynı klasörde tutun. Kontrol komutu aşağıdadır.'],
        ['10', 'Kontrol yalnız okur; yayın ayrı işlemdir. PDF dosyaları, soru/cevap ve öğrenci geçmişi değiştirilmez.'],
        ['11', 'Konu birleştirme onayı, kaynakların doğru etiketlendiği veya sayfa aralıklarının doğrulandığı anlamına gelmez.'],
        ['Kontrol', f'..\\.venv\\Scripts\\python.exe -X utf8 -B scripts/physics-topic-review.py validate "{directory / "fizik-konu-inceleme.xlsx"}"'],
        ['Tekil kaynaklar', json.dumps(snapshot['uniqueResources'], ensure_ascii=False)],
        ['Sayım', 'Eşleştirme satırlarındaki kaynak sayıları toplanmaz; aynı kaynak birden fazla satıra bağlı olabilir.'],
    ]
    for row in lines:
        info.append(row)
    for title, headers, rows in [('Temel konular', TOPIC_HEADERS, topics),
                                  ('Eşleştirmeler', MAP_HEADERS, mappings)]:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(65, max(16, len(str(col[0].value)) + 3))
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        decision = DataValidation(type='list', formula1='"BEKLIYOR,ONAY,ICERIK_INCELE"'
                                  if title == 'Eşleştirmeler' else '"BEKLIYOR,ONAY"')
        decision.errorTitle, decision.error = 'Geçersiz karar', 'Listeden karar seçin.'
        decision.showErrorMessage = True
        sheet.add_data_validation(decision)
        decision.add('G2:G501' if title == 'Eşleştirmeler' else 'E2:E501')
    workbook.defined_names.add(DefinedName('TemelKonuKodlari', attr_text="'Temel konular'!$A$2:$A$501"))
    targets = DataValidation(type='list', formula1='TemelKonuKodlari', allow_blank=True)
    workbook['Eşleştirmeler'].add_data_validation(targets)
    targets.add('E2:E501')
    exams = DataValidation(type='list', formula1='"TYT,AYT"')
    workbook['Temel konular'].add_data_validation(exams)
    exams.add('B2:B501')
    info.column_dimensions['A'].width = 25
    info.column_dimensions['B'].width = 125
    for row in info:
        for cell in row:
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        info.row_dimensions[row[0].row].height = 42
    destination = directory / 'fizik-konu-inceleme.xlsx'
    workbook.save(destination)
    print(json.dumps({'workbook': str(destination), 'topics': len(topics), 'categories': len(mappings),
                      'suggested': sum(bool(row[4]) for row in mappings),
                      'uniqueResources': snapshot['uniqueResources'], 'published': False}, ensure_ascii=False))


def sheet_rows(workbook, name, headers):
    sheet = workbook[name]
    if list(next(sheet.values)) != headers:
        raise ValueError(f'{name}: sütun başlıkları değişmiş.')
    for cells in list(sheet.rows)[1:]:
        if any(cell.data_type == 'f' for cell in cells):
            raise ValueError(f'{name}: formül yerine doğrudan değer girin.')
    return [[str(value).strip() if value is not None else '' for value in row]
            for row in list(sheet.values)[1:] if any(value is not None for value in row)]


def validate(path):
    baseline = json.loads((path.parent / 'baseline.json').read_text(encoding='utf-8'))
    current = read_snapshot()
    if current['runKey'] != baseline['runKey'] or current['taxonomySha256'] != baseline['taxonomySha256']:
        raise ValueError('Kaynak sürümü değişmiş; öğretmen kararlarını yeni sürümle uzlaştırın.')
    workbook = load_workbook(path, data_only=False)
    topic_rows = sheet_rows(workbook, 'Temel konular', TOPIC_HEADERS)
    mappings = sheet_rows(workbook, 'Eşleştirmeler', MAP_HEADERS)
    errors, pending, per_item = [], [], []
    topics, names = {}, set()
    for code, exam, name, note, decision in topic_rows:
        if not re.fullmatch(r'F\d{3,}', code) or code in topics or exam not in {'TYT', 'AYT'} or not normal(name):
            errors.append(f'Geçersiz/tekrarlı temel konu: {code} {name}')
        if (exam, normal(name)) in names:
            errors.append(f'Aynı anlamı taşıyabilecek tekrarlı ad: {exam} {name}')
        names.add((exam, normal(name)))
        topics[code] = {'exam': exam, 'name': name, 'scope': note, 'decision': decision}
        if decision not in {'ONAY', 'BEKLIYOR', ''}:
            errors.append(f'Geçersiz temel konu kararı: {code}')
        if decision != 'ONAY':
            pending.append(code)
    if not topics:
        errors.append('Temel konu listesi boş.')
    original = {row[0]: [str(value) for value in row] for row in baseline['mappingRows']}
    seen, approved, used = set(), [], set()
    for row in mappings:
        key, exam, _, _, target, _, decision, note, *_ = row
        if key not in original or key in seen:
            errors.append(f'Bilinmeyen/tekrarlı eski kimlik: {key}')
            continue
        seen.add(key)
        if any(row[i] != original[key][i] for i in [0, 1, 2, 3, 5, 8, 9, 10, 11, 12]):
            errors.append(f'Kaynak sütunu değişmiş: {key}')
        if decision not in {'ONAY', 'BEKLIYOR', '', 'ICERIK_INCELE'}:
            errors.append(f'Geçersiz karar: {key}')
        if target and (target not in topics or topics[target]['exam'] != exam):
            errors.append(f'Geçersiz hedef veya sınav kapsamı: {key} -> {target}')
        if decision == 'ONAY':
            if not target:
                errors.append(f'Onaylı eşleşmede hedef boş: {key}')
            else:
                used.add(target)
                approved.append({'from': key, 'targetCode': target, 'note': note})
        elif decision == 'ICERIK_INCELE':
            if target or not note:
                errors.append(f'İçerik incelemesinde hedef boş ve gerekçe dolu olmalı: {key}')
            per_item.append(key)
        else:
            pending.append(key)
    if seen != original.keys():
        errors.append('330 eski kategorinin tamamı korunmalı; eşleştirme satırını silmeyin.')
    ready = not errors and not pending and not per_item
    report = {'version': 1, 'runKey': baseline['runKey'], 'taxonomySha256': baseline['taxonomySha256'],
              'workbookSha256': digest(path.read_bytes()), 'topics': topics, 'topicCount': len(topics),
              'approvedMappings': approved, 'itemReviewRequired': per_item, 'pending': pending,
              'unusedTopicCodes': sorted(topics.keys() - used), 'errors': errors,
              'migrationPlanReady': ready, 'published': False, 'databaseWrites': 0,
              'note': 'Kontrol sonucu içerik doğruluğu onayı veya uygulama/yayın değildir.'}
    output = path.with_suffix('.validation.json')
    write_json(output, report)
    print(json.dumps({'report': str(output), 'topics': len(topics), 'approved': len(approved),
                      'pending': len(pending), 'itemReviewRequired': len(per_item), 'errors': errors,
                      'migrationPlanReady': ready, 'published': False}, ensure_ascii=False))
    return 1 if errors else 0


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    sub.add_parser('export')
    check = sub.add_parser('validate')
    check.add_argument('workbook', type=Path)
    args = parser.parse_args()
    if args.command == 'export':
        export()
        return 0
    return validate(args.workbook.resolve())


if __name__ == '__main__':
    raise SystemExit(main())
